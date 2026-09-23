from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

def add_section_border(paragraph):
    """Add a bottom border to a paragraph (used for section headers)."""
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '4472C4')
    pBdr.append(bottom)
    pPr.append(pBdr)

def add_code_block(doc, code_text):
    """Add a styled code block paragraph."""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.3)
    run = p.add_run(code_text)
    run.font.name = 'Courier New'
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    p.style = doc.styles['No Spacing']
    shading = OxmlElement('w:shd')
    shading.set(qn('w:val'), 'clear')
    shading.set(qn('w:color'), 'auto')
    shading.set(qn('w:fill'), 'F2F2F2')
    p._p.get_or_add_pPr().append(shading)
    return p

def set_page_margins(doc, top=1.0, bottom=1.0, left=1.2, right=1.0):
    section = doc.sections[0]
    section.top_margin = Inches(top)
    section.bottom_margin = Inches(bottom)
    section.left_margin = Inches(left)
    section.right_margin = Inches(right)

def create_report():
    doc = Document()
    set_page_margins(doc)

    # ── COVER PAGE ──────────────────────────────────────────────────
    doc.add_paragraph()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = title.add_run("BCSE402L – Big Data Analytics")
    tr.bold = True
    tr.font.size = Pt(14)

    doc.add_paragraph()
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sub.add_run("Digital Assignment 2 (DA2) – Progress Report")
    sr.bold = True
    sr.font.size = Pt(18)
    sr.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

    doc.add_paragraph()
    ptitle = doc.add_paragraph()
    ptitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ptr = ptitle.add_run("Predictive Shipment Delay Analytics\nfor Logistics Decision Support")
    ptr.italic = True
    ptr.font.size = Pt(13)

    doc.add_paragraph()
    info = doc.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info.add_run("Team Members\n").bold = True
    info.add_run(
        "Pranav Satish Khadse  –  23BLC1237\n"
        "Abhinav Kumar Singh   –  23BLC1209\n"
        "Rishi Jain            –  23BLC1242\n"
        "Anmol Kapil           –  23BLC1229\n"
    )
    info.add_run("\nDate: September 2026")

    doc.add_page_break()

    # ── 1. APPLICATION OF CONCEPTS ──────────────────────────────────
    h = doc.add_heading('1. Application of Big Data Concepts', level=1)
    add_section_border(h)

    doc.add_paragraph(
        "For this project we applied three core Big Data techniques: HDFS for distributed storage, "
        "PySpark for distributed data processing, and Spark SQL for analytical querying."
    )

    doc.add_heading('1.1  Hadoop HDFS – Dataset Storage', level=2)
    doc.add_paragraph(
        "We installed Hadoop 3.3.6 natively on Windows and configured a single-node cluster with a NameNode and DataNode. "
        "The raw dataset (180,519 records, ~96 MB) was uploaded from the local filesystem into the HDFS virtual directory "
        "'/user/logistics/project_data/' using the CLI command below. "
        "This ensures data is managed by Hadoop rather than the OS filesystem, which is the first step toward distributed storage."
    )
    add_code_block(doc, "hdfs dfs -mkdir -p /user/logistics/project_data")
    add_code_block(doc, "hdfs dfs -put dataset/DataCoSupplyChainDataset.csv /user/logistics/project_data/")
    add_code_block(doc, "hdfs dfs -ls /user/logistics/project_data/")
    doc.add_paragraph("[ Insert screenshot of hdfs dfs -ls output here ]").italic = True

    doc.add_heading('1.2  PySpark – Distributed Processing & Transformation', level=2)
    doc.add_paragraph(
        "We initialized a PySpark SparkSession and read the dataset directly from HDFS. "
        "Spark distributes the read and subsequent operations across its worker threads automatically. "
        "We selected 8 relevant columns and removed any rows with null values in those fields."
    )
    add_code_block(doc, 'spark = SparkSession.builder.appName("DA2_Shipment_Delay_Analytics").getOrCreate()')
    add_code_block(doc, 'df = spark.read.option("header","true").option("encoding","iso-8859-1")')
    add_code_block(doc, '           .csv("hdfs://localhost:9000/user/logistics/project_data/DataCoSupplyChainDataset.csv")')
    add_code_block(doc, 'df_clean = df.select(selected_cols).dropna()')

    doc.add_heading('1.3  Spark SQL – Exploratory Data Analysis', level=2)
    doc.add_paragraph(
        "We registered the cleaned DataFrame as a temporary SQL view and ran three analytical queries "
        "to understand the delay patterns in the dataset:"
    )
    add_code_block(doc, 'df_clean.createOrReplaceTempView("shipments")')
    add_code_block(doc, '-- Query 1: Delay distribution')
    add_code_block(doc, 'SELECT Late_delivery_risk, COUNT(*) as count FROM shipments GROUP BY Late_delivery_risk')
    add_code_block(doc, '')
    add_code_block(doc, '-- Query 2: Delay rate by shipping mode')
    add_code_block(doc, 'SELECT `Shipping Mode`, AVG(Late_delivery_risk) as delay_rate, COUNT(*) as total_orders')
    add_code_block(doc, 'FROM shipments GROUP BY `Shipping Mode` ORDER BY delay_rate DESC')
    add_code_block(doc, '')
    add_code_block(doc, '-- Query 3: Top 10 regions by delay rate')
    add_code_block(doc, 'SELECT `Order Region`, AVG(Late_delivery_risk) as delay_rate')
    add_code_block(doc, 'FROM shipments GROUP BY `Order Region` ORDER BY delay_rate DESC LIMIT 10')

    doc.add_page_break()

    # ── 2. PROJECT PROGRESS ─────────────────────────────────────────
    h = doc.add_heading('2. Project Progress Summary', level=1)
    add_section_border(h)

    doc.add_paragraph(
        "The following tasks have been completed as part of DA2:"
    )
    items = [
        "Dataset obtained from Kaggle (DataCo Smart Supply Chain – 180,519 records, 53 columns).",
        "Local Hadoop 3.3.6 single-node cluster set up on Windows with NameNode formatted and DataNode running.",
        "Dataset uploaded to HDFS under /user/logistics/project_data/.",
        "PySpark pipeline implemented: reads CSV from HDFS, selects features, drops nulls.",
        "Three Spark SQL EDA queries written and executed; output saved as bar chart plots.",
        "Baseline Logistic Regression model trained (80/20 split, random_state=42).",
        "Model evaluated on test set: Accuracy 69.67%, ROC-AUC 74.09%.",
    ]
    for item in items:
        p = doc.add_paragraph(item, style='List Bullet')

    doc.add_paragraph(
        "\nAll results shown in this report are directly taken from the actual pipeline execution on our local machine."
    )
    doc.add_paragraph("[ Insert screenshot of full pipeline terminal output (Steps 1–6) here ]").italic = True

    doc.add_page_break()

    # ── 3. METHODS & TOOLS ──────────────────────────────────────────
    h = doc.add_heading('3. Methods & Tools Used', level=1)
    add_section_border(h)

    tools = [
        ("Hadoop 3.3.6 (HDFS)", "Used for distributed storage of the raw dataset. The NameNode manages metadata while the DataNode stores the actual data blocks. Chosen because it is the industry-standard Big Data storage layer."),
        ("Apache PySpark 4.2.0", "Used to read and process data from HDFS in a distributed manner. PySpark's DataFrame API allows SQL-like transformations with lazy evaluation and parallel execution."),
        ("Spark SQL", "Used for running SQL queries directly on distributed DataFrames via a temporary view, avoiding the need to move data to a relational database."),
        ("Scikit-Learn 1.x", "Used for the baseline Logistic Regression model. The OneHotEncoder handled categorical columns and the Pipeline object combined preprocessing and training into a single step."),
        ("Matplotlib / Seaborn", "Used to visualize EDA query results as bar charts and horizontal bar plots."),
        ("Python-docx", "Used to programmatically generate this report."),
        ("Java 17 (Eclipse Temurin)", "Required runtime for both Hadoop and PySpark. Installed to JAVA_HOME=C:\\java."),
    ]
    for tool, desc in tools:
        p = doc.add_paragraph()
        p.add_run(f"{tool}: ").bold = True
        p.add_run(desc)

    doc.add_page_break()

    # ── 4. TEAM PARTICIPATION ───────────────────────────────────────
    h = doc.add_heading('4. Team Participation', level=1)
    add_section_border(h)

    contributions = [
        ("Pranav Satish Khadse (23BLC1237)", "Hadoop setup, NameNode formatting, HDFS configuration (core-site.xml, hdfs-site.xml). Dataset upload to HDFS via CLI."),
        ("Abhinav Kumar Singh (23BLC1209)", "PySpark session initialization, CSV ingestion from HDFS, DataFrame column selection and null-dropping transformations."),
        ("Rishi Jain (23BLC1242)", "Spark SQL query design and execution for all three EDA queries. Saved results to Pandas for visualization."),
        ("Anmol Kapil (23BLC1229)", "Machine learning pipeline: One-Hot Encoding, train/test split, Logistic Regression training, metric evaluation, and report preparation."),
    ]
    for name, work in contributions:
        p = doc.add_paragraph()
        p.add_run(f"{name}\n").bold = True
        p.add_run(work)
        doc.add_paragraph()

    doc.add_page_break()

    # ── 5. EXPLANATION OF WORK ──────────────────────────────────────
    h = doc.add_heading('5. Detailed Explanation of Work Completed', level=1)
    add_section_border(h)

    doc.add_heading('5.1  Pipeline Workflow', level=2)
    doc.add_paragraph(
        "The entire pipeline runs through a single Python file (run_da2_project.py). The flow is as follows:\n"
        "1. SparkSession is started and connected to the local Hadoop cluster.\n"
        "2. The CSV is read from HDFS (180,519 rows loaded).\n"
        "3. Eight columns are selected; rows with any null value are dropped. All 180,519 rows were retained "
        "since the selected columns had no missing values.\n"
        "4. A Spark SQL temporary view 'shipments' is registered and three queries are executed.\n"
        "5. Query results are converted to Pandas and saved as .png charts.\n"
        "6. The cleaned Spark DataFrame is collected to Pandas; features are One-Hot Encoded and the "
        "Logistic Regression model is trained and evaluated."
    )

    doc.add_heading('5.2  EDA Results', level=2)
    doc.add_paragraph("Query 1 – Delay Distribution: Out of 180,519 shipments, 98,977 (54.8%) were delayed and 81,542 (45.2%) were on time, indicating a class imbalance that will need to be addressed in DA3.")
    if os.path.exists("eda_q1_distribution.png"):
        doc.add_picture("eda_q1_distribution.png", width=Inches(3.8))
    doc.add_paragraph()

    doc.add_paragraph("Query 2 – Delay Rate by Shipping Mode: First Class had by far the highest delay rate (95.3%), while Standard Class had the lowest (38.1%). This is a key feature for the predictive model.")
    if os.path.exists("eda_q2_shipping_mode.png"):
        doc.add_picture("eda_q2_shipping_mode.png", width=Inches(4.5))
    doc.add_paragraph()

    doc.add_paragraph("Query 3 – Top Regions by Delay Rate: Central Africa (57.9%), South Asia (56.3%), and East Africa (55.9%) had the highest delay rates, suggesting geographical factors are important predictors.")
    if os.path.exists("eda_q3_region.png"):
        doc.add_picture("eda_q3_region.png", width=Inches(5.0))
    doc.add_paragraph()

    doc.add_heading('5.3  Baseline Model Results', level=2)
    doc.add_paragraph(
        "The Logistic Regression model was evaluated on 36,104 test samples:"
    )
    metrics = [
        ("Accuracy",  "69.67%"),
        ("Precision", "84.30%"),
        ("Recall",    "54.91%"),
        ("F1-Score",  "66.50%"),
        ("ROC-AUC",   "74.09%"),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.style = 'Light Shading Accent 1'
    hdr = table.rows[0].cells
    hdr[0].text = "Metric"
    hdr[1].text = "Value"
    for m, v in metrics:
        row = table.add_row().cells
        row[0].text = m
        row[1].text = v

    doc.add_paragraph()
    doc.add_paragraph(
        "The model achieves good precision (84.3%), meaning when it predicts a delay, it is correct most of the time. "
        "Recall is lower (54.9%), indicating it misses some actual delays. This is expected from a baseline model "
        "using only categorical features. In DA3, we plan to incorporate numerical features and try Random Forest "
        "or XGBoost to improve recall and overall F1 score."
    )

    doc.save("DA2_Report_Final.docx")
    print("Report saved: DA2_Report_Final.docx")

if __name__ == "__main__":
    create_report()
